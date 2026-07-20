import sys, os
sys.stdout.reconfigure(encoding='utf-8')

try:
    from docx import Document
except ImportError:
    os.system('pip install python-docx -q')
    from docx import Document

path = r'd:\iCode\cursor\d_hos_backend_cc_20251230\微爱基金会服务采购协议20251201(1).docx'
doc = Document(path)

for i, para in enumerate(doc.paragraphs):
    if para.text.strip():
        print(para.text)

# 也打印表格内容
for table in doc.tables:
    for row in table.rows:
        cells = [c.text.strip() for c in row.cells if c.text.strip()]
        if cells:
            print(' | '.join(cells))
