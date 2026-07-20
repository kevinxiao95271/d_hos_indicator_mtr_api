#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Evaluation indicator system - MD to DOCX converter
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
import re
from pathlib import Path

def set_font(run, size=11, bold=False, name='宋体'):
    """设置字体"""
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.bold = bold

def md_to_docx(md_file, docx_file):
    """将Markdown文件转换为DOCX"""

    # 读取MD文件
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()

    # 创建Document
    doc = Document()

    # 设置样式
    style = doc.styles['Normal']
    style.font.name = '宋体'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    style.font.size = Pt(11)

    # 分行处理
    lines = content.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i]

        # 跳过空行
        if not line.strip():
            i += 1
            continue

        # 处理标题
        if line.startswith('# '):
            p = doc.add_heading(line[2:].strip(), level=1)
            for run in p.runs:
                set_font(run, size=16, bold=True)
        elif line.startswith('## '):
            p = doc.add_heading(line[3:].strip(), level=2)
            for run in p.runs:
                set_font(run, size=14, bold=True)
        elif line.startswith('### '):
            p = doc.add_heading(line[4:].strip(), level=3)
            for run in p.runs:
                set_font(run, size=12, bold=True)

        # 处理表格
        elif line.startswith('|'):
            table_lines = [line]
            i += 1
            # 收集表格行
            while i < len(lines) and lines[i].startswith('|'):
                table_lines.append(lines[i])
                i += 1

            # 解析表格
            rows = []
            for table_line in table_lines:
                cells = [cell.strip() for cell in table_line.split('|')[1:-1]]
                if cells:
                    rows.append(cells)

            # 跳过分隔符行
            if len(rows) > 1:
                rows = [rows[0]] + [r for r in rows[2:] if r]

                # 创建表格
                if rows:
                    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                    table.style = 'Table Grid'

                    # 设置表头
                    for cell_idx, cell_text in enumerate(rows[0]):
                        cell = table.rows[0].cells[cell_idx]
                        cell.text = cell_text
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                set_font(run, size=10, bold=True)

                    # 填充数据
                    for row_idx in range(1, len(rows)):
                        for cell_idx, cell_text in enumerate(rows[row_idx]):
                            cell = table.rows[row_idx].cells[cell_idx]
                            cell.text = cell_text
                            for paragraph in cell.paragraphs:
                                for run in paragraph.runs:
                                    set_font(run, size=10)
            continue

        # 处理代码块
        elif line.startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                code_lines.append(lines[i])
                i += 1

            p = doc.add_paragraph()
            for code_line in code_lines:
                r = p.add_run(code_line + '\n')
                set_font(r, size=9, name='Courier New')
                r.font.color.rgb = RGBColor(100, 100, 100)
            i += 1
            continue

        # 处理列表
        elif line.startswith('- '):
            p = doc.add_paragraph(line[2:].strip(), style='List Bullet')
            for run in p.runs:
                set_font(run, size=11)

        # 处理普通段落
        else:
            p = doc.add_paragraph(line.strip())
            for run in p.runs:
                set_font(run, size=11)

        i += 1

    # 保存
    doc.save(docx_file)
    print('[OK] converted: {} to {}'.format(md_file, docx_file))

if __name__ == '__main__':
    base_path = Path('d:/AiCode/cla_vs_fchat/d_hos_backend_cc_20251230')

    # 三份文档转换
    files = [
        ('评审指标系统设计分析报告.md', 'evaluation_indicator_deliverables/1_评审指标系统设计分析报告.docx'),
        ('评审指标系统实施指南.md', 'evaluation_indicator_deliverables/2_评审指标系统实施指南.docx'),
        ('评审指标系统-最小粒子层实施成果总结.md', 'evaluation_indicator_deliverables/3_评审指标系统成果总结.docx'),
    ]

    for md_name, docx_name in files:
        md_file = base_path / md_name
        docx_file = base_path / docx_name

        if md_file.exists():
            md_to_docx(str(md_file), str(docx_file))
        else:
            print('[ERROR] File not found: {}'.format(md_file))

    print('\\nAll conversions completed!')
